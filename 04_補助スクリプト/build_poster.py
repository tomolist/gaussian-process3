# -*- coding: utf-8 -*-
"""A0 縦のポスター（PowerPoint）を作る：05_…/学習結果報告書作成_261015/ポスター/報告書Bから作成.pptx

内容は、学習結果報告書作成_261015 の報告書（過去報告書A・報告書B）と、見本のポスター（ポスター/報告書Aから作成2.pdf）の書式に合わせる。
読み手：プラスチック成形の知識はあるが、機械学習は知らない人。
図は make_poster_figs.py が作ったもの（ポスター/図/）と、報告書B の図3・図4（図/Fig3.png・Fig4.png）。
文字はテキストボックス、表は PowerPoint の表、図は画像として置くので、PowerPoint で後から直せる。

使い方（リポジトリ直下で）:
  uv run python 04_補助スクリプト/make_poster_figs.py   （図を作る。先に 1 回）
  uv run python 04_補助スクリプト/build_poster.py
"""
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

REPO = "."
REPORT = os.path.join(REPO, "05_材料ごとの機械学習データと予測検証データ", "学習結果報告書作成_261015")
POSTER = os.path.join(REPORT, "ポスター")
FIG = os.path.join(POSTER, "図")
OUT = os.path.join(POSTER, "報告書Bから作成.pptx")

FONT = "BIZ UDPゴシック"
BLUE = "006699"          # 見出しの帯（見本と同じ）
FRAME = "2E75B6"         # 節の枠線（見本と同じ）
GREEN = "A9CE98"         # 小見出し（見本と同じ）
GREEN_LINE = "70AD47"
RED = "E00000"
LIGHT = "EEF6FB"         # 囲みの薄い青
ORANGE = "F8CBAD"
ORANGE_LINE = "C55A11"
SKY = "BDD7EE"

PAGE_W, PAGE_H = 841.0, 1189.0   # A0 縦 [mm]
M = 8.0                           # 外側の余白
GAP = 8.0                         # 列の間
CW = (PAGE_W - 2 * M - GAP) / 2   # 列の幅
XL, XR = M, M + CW + GAP          # 左右の列の左端
PAD = 5.0                         # 枠の内側の余白
IW = CW - 2 * PAD                 # 列の中身の幅
HEAD_H = 19.0                     # 節の見出しの帯の高さ
SUB_H = 11.0                      # 小見出しの高さ
BODY = 22                         # 本文の文字の大きさ [pt]
NOTE = 19                         # 注の文字の大きさ [pt]

PT_MM = 25.4 / 72.0
EMU_MM = 36000.0


def rgb(h):
    return RGBColor.from_string(h)


def est_width(text, size):
    """文字列の幅の目安 [mm]（全角 1 字 = 1 em、半角 = 0.55 em）"""
    text = re.sub(r"\[/?[rb]\]|[_^]\{|\}", "", text)
    return sum(1.0 if ord(c) > 0x2000 else 0.55 for c in text) * size * PT_MM


# ---------------------------------------------------------------- 文字の書き込み（簡単な書式：[r]赤[/r]、[b]太字[/b]、_{下付き}、^{上付き}）
TOKEN = re.compile(r"(\[r\]|\[/r\]|\[b\]|\[/b\]|_\{[^}]*\}|\^\{[^}]*\})")
MATH_FONT = "Cambria Math"        # γ̇（γ ＋ 上の点の結合文字）は、BIZ UDPゴシックでは点がずれるので、この字だけ数式用のフォントで書く
SYM = re.compile("(γ̇)")


def style_run(run, size, color, bold, baseline=None, font=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
    if baseline:
        rPr.set("baseline", str(baseline))


def add_runs(p, text, size, color="000000", bold=False):
    red = b = False
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok in ("[r]", "[/r]"):
            red = tok == "[r]"
            continue
        if tok in ("[b]", "[/b]"):
            b = tok == "[b]"
            continue
        base = None
        if tok.startswith("_{"):
            tok, base = tok[2:-1], -25000
        elif tok.startswith("^{"):
            tok, base = tok[2:-1], 30000
        for piece in SYM.split(tok):
            if not piece:
                continue
            r = p.add_run()
            r.text = piece
            style_run(r, size, RED if red else color, bold or b, base, font=MATH_FONT if SYM.fullmatch(piece) else FONT)


def fill_text(tf, text, size, color="000000", bold=False, align="left", spacing=1.0, after=0.0, hang=None):
    """text を段落（\n 区切り）に分けて書く。hang：行頭が「・①②③」の段落を、ぶら下げにする幅（em）"""
    al = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        p.line_spacing = spacing
        p.space_after = Pt(after)
        if hang and line[:1] in "・①②③④":
            pPr = p._p.get_or_add_pPr()
            w = int(hang * size * 12700)
            pPr.set("marL", str(w))
            pPr.set("indent", str(-w))
        add_runs(p, line, size, color, bold)


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.5, radius=None):
    s = slide.shapes.add_shape(kind, Mm(x), Mm(y), Mm(w), Mm(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    if radius is not None:
        s.adjustments[0] = radius
    return s


def text_in(s, text, size, color="000000", bold=False, align="left", anchor="top", margin=1.5, spacing=1.0, after=0.0, hang=None):
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Mm(margin))
    fill_text(tf, text, size, color, bold, align, spacing, after, hang)
    return s


def tbox(slide, x, y, w, h, text, size=BODY, **kw):
    s = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    return text_in(s, text, size, **kw)


def box(slide, x, y, w, h, text, size=BODY, fill="FFFFFF", line=BLUE, lw=2.0, radius=0.08, **kw):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius)
    kw.setdefault("margin", 2.5)
    return text_in(s, text, size, **kw)


def pic(slide, name, x, y, w=None, h=None, folder=FIG):
    path = os.path.join(folder, name)
    assert os.path.exists(path), f"図がない：{path}（先に make_poster_figs.py を実行）"
    p = slide.shapes.add_picture(path, Mm(x), Mm(y), Mm(w) if w else None, Mm(h) if h else None)
    return p.width / EMU_MM, p.height / EMU_MM


def header(slide, x, y, w, title):
    s = shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, HEAD_H, fill=BLUE)
    text_in(s, title, 46, color="FFFFFF", bold=True, anchor="middle", margin=4.0)
    return y + HEAD_H


def frame(slide, x, y0, w, y1):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y0, w, y1 - y0, line=FRAME, lw=2.0)


def sub(slide, x, y, title, size=26):
    w = est_width(title, size) + 8.0
    s = shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, SUB_H, fill=GREEN, line=GREEN_LINE, lw=1.25)
    text_in(s, title, size, bold=True, anchor="middle", margin=3.0)
    return y + SUB_H + 2.5


def arrow(slide, kind, x, y, w, h):
    shape(slide, kind, x, y, w, h, fill="FFFFFF", line=BLUE, lw=2.25)


def row_height(size, lines=1):
    """表の 1 行の高さ [mm]（PowerPoint の行の高さ：文字の大きさ × 1.2 × 行数 ＋ 上下の余白 0.8 mm × 2）"""
    return size * PT_MM * 1.2 * lines + 1.6 + 0.4


def table(slide, x, y, widths, rows, size=NOTE, row_h=None, head_fill=BLUE, first_col_fill="EDEDED", merges=(), aligns=None, lines=None):
    """rows：文字列の 2 次元リスト（1 行目は見出し）。merges：((行, 列から, 列まで), ...) を横に結合。lines：行ごとの文字の行数（既定 1）。
    返り値：表の下端 [mm]"""
    n, m = len(rows), len(widths)
    lines = lines or [1] * n
    hs = [row_height(size, k) if row_h is None else row_h * k for k in lines]
    gf = slide.shapes.add_table(n, m, Mm(x), Mm(y), Mm(sum(widths)), Mm(sum(hs)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"           # スタイルなし、表のグリッド線
    for j, w in enumerate(widths):
        tbl.columns[j].width = Mm(w)
    for i in range(n):
        tbl.rows[i].height = Mm(hs[i])
    for (i, j0, j1) in merges:
        tbl.cell(i, j0).merge(tbl.cell(i, j1))
    for i, row in enumerate(rows):
        for j, t in enumerate(row):
            c = tbl.cell(i, j)
            if c.is_spanned:
                continue
            c.margin_left = c.margin_right = Mm(1.5)
            c.margin_top = c.margin_bottom = Mm(0.8)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i == 0 and head_fill:
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(head_fill)
            elif j == 0 and first_col_fill:
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(first_col_fill)
            else:
                c.fill.solid()
                c.fill.fore_color.rgb = rgb("FFFFFF")
            tf = c.text_frame
            tf.word_wrap = True
            al = (aligns[j] if aligns else ("left" if j == 0 else "center"))
            fill_text(tf, t, size, color="FFFFFF" if (i == 0 and head_fill) else "000000", bold=(i == 0 or j == 0), align=al)
    return y + sum(hs)


# ---------------------------------------------------------------- ポスター
def title_band(s):
    h = 64.0
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, PAGE_W, h, fill=BLUE)
    tbox(s, M, 3.5, PAGE_W - 2 * M, 24.0, "金型内センサデータを用いたデータ駆動型樹脂粘度モデルの開発Ⅲ", 60, color="FFFFFF", bold=True, align="center")
    tbox(s, M, 27.0, PAGE_W - 2 * M, 11.0, "Development of Data-Driven Melt Viscosity Model using Data Measured by In-Mold Sensors Ⅲ", 26,
         color="FFFFFF", align="center")
    tbox(s, M, 38.5, PAGE_W - 2 * M, 12.0, "〔著者名を記入〕", 28, color="FFFFFF", bold=True, align="center")
    tbox(s, M, 51.0, PAGE_W - 2 * M, 10.0, "〔所属を記入〕", 22, color="FFFFFF", align="center")
    return h + 5.0


def left_column(s, y):
    x, xi = XL, XL + PAD
    # ---------------- 1. 研究概要
    y0 = y
    y = header(s, x, y, CW, "1. 研究概要") + 2.5
    t = ("　射出成形 CAE で用いられる樹脂粘度モデル^{1)} には、樹脂の製造ロット間の分子量分布のばらつきや、成形工程の熱履歴による"
         "ばらつきが考慮されていない。本研究では、金型内の樹脂圧力と温度を計測できる計測金型と機械学習を用いて、[r]樹脂粘度の予測値に"
         "加えて、予測のばらつきの範囲も同時に求められる[/r]、新たな[r]データ駆動型樹脂粘度モデル[/r]の構築を目的としている。")
    tbox(s, xi, y, IW, 40.0, t, BODY, spacing=1.0)
    y += 41.0
    pts = ("[b]本報のポイント（前報からの発展）[/b]\n"
           "① 3 種類の樹脂（PP・粉砕材混合 PP・GPPS）でモデルを構築\n"
           "② Cross-WLF モデルと同様の[r]温度シフトファクター[/r]をモデルに組み込み\n"
           "③ [r]学習に用いていない射出率の条件[/r]（学習範囲の外）で予測精度を検証")
    box(s, xi, y, IW, 46.0, pts, BODY, fill=LIGHT, hang=1.0, spacing=1.0)
    y += 49.0
    frame(s, x, y0, CW, y)
    y += 5.0

    # ---------------- 2. 計測金型およびセンサ
    y0 = y
    y = header(s, x, y, CW, "2. 計測金型およびセンサ") + 3.0
    w1, h1 = pic(s, "見本_計測金型.png", xi, y, h=100.0)
    w2, h2 = pic(s, "見本_キャビティ入れ子.png", xi + IW - 100.0 * 1908 / 1155, y, h=100.0)
    gx0, gx1 = xi + w1 + 4.0, xi + IW - w2 - 4.0                       # 2 枚の写真の間の空き
    box(s, gx0, y + 30.0, gx1 - gx0, 40.0, "キャビティブロック（入れ子）の交換で、厚み 0.5・1.0・1.5・2.0 mm のキャビティで計測が可能", 20,
        anchor="middle", align="center", margin=2.0)
    y += 102.0
    wd = 140.0
    w3, h3 = pic(s, "見本_センサの設置位置.png", xi, y, w=IW - wd - 6.0)
    sens = ("①②：水晶圧電式樹脂圧力センサ（6157CAUMRKF0.4、日本キスラー合同会社）\n"
            "②：赤外線放射樹脂温度センサ（EPSSZLφ4、双葉電子工業㈱）\n"
            "専用の固定側キャビティブロックの交換で、圧力と温度を別々に計測")
    tbox(s, xi, y + h3 + 1.0, IW - wd - 6.0, 30.0, sens, NOTE, hang=None, spacing=1.0)
    w4, h4 = pic(s, "見本_成形品の形状.png", xi + IW - wd, y, w=wd)
    tbox(s, xi + IW - wd, y + h4, wd, 9.0, "キャビティ形状・寸法（単位：mm）", NOTE, align="center")
    y += max(h3 + 31.0, h4 + 9.0) + 2.0
    frame(s, x, y0, CW, y)
    y += 5.0

    # ---------------- 3. 成形条件
    y0 = y
    y = header(s, x, y, CW, "3. 成形条件") + 3.0
    rows = [["", "PP", "粉砕材混合 PP", "GPPS"],
            ["樹脂", "プライムポリプロ J707EG（㈱プライムポリマー）", "同 PP に、一次成形品の粉砕材を 25 wt% 混合", "ディックスチレン CR-2500（DIC㈱）"],
            ["加熱シリンダ温度", "190・210・230 ℃", "190・210・230 ℃", "200・220・240 ℃"],
            ["キャビティ厚み", "0.5・1.0・1.5・2.0 mm", "0.5・1.0・1.5・2.0 mm", "1.0・1.5・2.0 mm"],
            ["射出率", "1.6〜170 cm^{3}/s の 12 水準（3 材料共通）", "", ""],
            ["保持圧力", "0 MPa", "", ""],
            ["成形機", "全電動式射出成形機 ROBOSHOT α-50iAP（ファナック㈱、スクリュ直径 26 mm）", "", ""]]
    y = table(s, xi, y, [78.0, 107.0, 107.0, IW - 292.0], rows, NOTE, merges=[(4, 1, 3), (5, 1, 3), (6, 1, 3)],
              lines=[1, 2, 1, 1, 1, 1, 1]) + 1.5                         # 樹脂の行は 2 行に折り返す
    tbox(s, xi, y, IW, 9.0, "※薄肉・低射出率の一部の条件は成形できないため、全組合せはそろっていない", NOTE)
    y += 10.0
    frame(s, x, y0, CW, y)
    y += 5.0

    # ---------------- 4. 実験結果（左の列）
    y0 = y
    y = header(s, x, y, CW, "4. 実験結果") + 3.0
    y = sub(s, xi, y, "樹脂圧力と樹脂温度の経時変化の一例")
    wp = 196.0
    w5, h5 = pic(s, "Fig3.png", xi, y, w=wp, folder=os.path.join(REPORT, "図"))
    w6, h6 = pic(s, "見本_樹脂温度の経時変化.png", xi + IW - h5 * 2416 / 1680, y, h=h5)
    y += h5 + 1.0
    tbox(s, xi, y, IW, 8.5, "PP、キャビティ厚み 1.0 mm、加熱シリンダ温度条件Ⅱ、射出率 42.5 cm^{3}/s　（左：樹脂圧力、右：樹脂温度）", NOTE, align="center")
    y += 10.0
    ex = ("・樹脂圧力は、センサ位置①→②の順に立ち上がる\n"
          "・Δt：①と②の圧力の立ち上がり時間の差 → 樹脂流量 Q を算出\n"
          "・ΔP_{1}：②の立ち上がり時刻における①と②の圧力差 → 樹脂粘度 η を算出\n"
          "・ΔP_{2}：②の最大値時における①と②の圧力差 → [r]機械学習の入力[/r]\n"
          "・樹脂温度：フローフロントが②に到達した時点で最高値（金型内樹脂最高温度）→ [r]機械学習の入力[/r]")
    box(s, xi, y, IW, 48.0, ex, 20, hang=1.0, spacing=1.0)
    y += 50.0
    frame(s, x, y0, CW, y)
    return y


def right_column(s, y, min_bottom=0.0):
    """min_bottom：最後の枠の下端を、ここまで延ばす（左の列の下端にそろえる）"""
    x, xi = XR, XR + PAD
    # ---------------- 4. 実験結果（続き）
    y0 = y
    y += 3.0
    y = sub(s, xi, y, "樹脂粘度とせん断速度の算出（ニュートン流体の計算式^{2)}）")
    wc, hc = pic(s, "見本_キャビティの模式図.png", xi, y, w=200.0)
    wf, hf = pic(s, "算出式.png", xi + IW - 190.0, y, w=190.0)
    y += max(hc, hf) + 0.5
    tbox(s, xi, y, IW, 8.5, "L：センサ間距離（17.5 mm）、B：流路幅、H：流路高さ（キャビティ厚み）、Q：樹脂流量、V：流速", NOTE)
    y += 10.5
    y = sub(s, xi, y, "樹脂粘度とせん断速度の関係")
    w7, h7 = pic(s, "Fig4.png", xi, y, w=IW, folder=os.path.join(REPORT, "図"))
    y += h7 + 2.0
    bw = 150.0
    box(s, xi, y, bw, 25.0, "せん断速度の増加、加熱シリンダ温度の上昇に伴い、粘度が低下", 20, anchor="middle")
    box(s, xi + bw + 3.0, y, bw, 25.0, "高せん断速度側のプロットは、キャピラリーレオメータの計測データとおおむね一致", 20, anchor="middle")
    arrow(s, MSO_SHAPE.RIGHT_ARROW, xi + 2 * bw + 6.0, y + 5.0, 14.0, 15.0)
    box(s, xi + 2 * bw + 23.0, y, IW - 2 * bw - 23.0, 25.0, "[r]学習データとして使用可能[/r]", 20, anchor="middle", align="center")
    y += 28.0
    frame(s, x, y0, CW, y)
    y += 5.0

    # ---------------- 5. 機械学習（ガウス過程回帰）
    y0 = y
    y = header(s, x, y, CW, "5. 機械学習（ガウス過程回帰）") + 3.0
    y = sub(s, xi, y, "ガウス過程回帰とは")
    t = ("　粘度式の形（Cross 式など）をあらかじめ決めずに、入力（成形・計測条件）と出力（粘度）の関係をデータから学習する手法^{3)}。"
         "[r]予測値とともに、そのばらつきの範囲（予測区間）[/r]を求めることができる。")
    tbox(s, xi, y, IW, 19.0, t, BODY)
    y += 19.5
    wg, hg = pic(s, "GPR概念図.png", xi, y, w=195.0)
    bx, bw = xi + 199.0, IW - 199.0
    box(s, bx, y + 1.0, bw, 30.0, "[b]従来の粘度式[/b]（Cross-WLF モデルなど）\n式の形を決め、係数をデータに合わせる\n→ 予測値のみ", 19,
        line="7F7F7F", spacing=1.0)
    box(s, bx, y + 34.0, bw, 30.0, "[b]ガウス過程回帰[/b]\n式の形を決めずに、データから関係を学ぶ\n→ [r]予測値 ＋ 予測のばらつきの範囲[/r]", 19, spacing=1.0)
    tbox(s, bx, y + 67.0, bw, 40.0, "[b]95% 予測区間[/b]：同じ条件で計測を繰り返したとき、計測値の 95% が入ると予測される範囲。データが少ない所ほど広くなる", NOTE)
    y += hg + 2.0

    y = sub(s, xi, y, "本モデルの入力と出力")
    hb = 46.0
    w_in, w_ar, w_md = 128.0, 13.0, 112.0
    w_out = IW - w_in - w_md - 2 * w_ar - 12.0
    box(s, xi, y, w_in, hb, "[b]入力（説明変数）[/b]\n・せん断速度 γ̇\n・樹脂圧力差 ΔP_{2}\n・金型内樹脂最高温度\n・加熱シリンダ温度", 20, fill=LIGHT,
        hang=1.0, spacing=0.95)
    arrow(s, MSO_SHAPE.RIGHT_ARROW, xi + w_in + 3.0, y + hb / 2 - 7.0, w_ar, 14.0)
    m = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, xi + w_in + w_ar + 6.0, y, w_md, hb, fill=BLUE, radius=0.08)
    text_in(m, "ガウス過程回帰\nモデル\n（温度シフトファクター付き）", 21, color="FFFFFF", bold=True, align="center", anchor="middle")
    arrow(s, MSO_SHAPE.RIGHT_ARROW, xi + w_in + w_ar + w_md + 9.0, y + hb / 2 - 7.0, w_ar, 14.0)
    box(s, xi + IW - w_out, y, w_out, hb, "[b]出力（目的変数）[/b]\n樹脂粘度 η の\n[r]予測値と 95% 予測区間[/r]", 20, fill=LIGHT,
        align="center", anchor="middle", spacing=1.0)
    y += hb + 2.0
    notes = ("・η・γ̇・ΔP_{2} は対数に変換。γ̇・ΔP_{2}・金型内樹脂最高温度には、同じ成形条件のショットの平均値を使用（ショットごとの値では、"
             "計測値の小さなゆらぎに予測曲線が過敏に反応し、学習範囲外の予測精度が低下したため）\n"
             "・カーネル関数^{4)}（データ同士の似ている度合いを表す関数）：入力ごとに影響の大きさを自動で調整する ARD カーネルに、"
             "学習範囲の外で予測値が平均値に戻るのを防ぐ直線的な成分と、ショット間のばらつきを表す成分を加えた")
    tbox(s, xi, y, IW, 36.0, notes, NOTE, hang=1.0, spacing=1.0)
    y += 37.0

    y = sub(s, xi, y, "学習データと検証データ（射出率で分ける）")
    rates = ["1.6", "2.7", "5.3", "10.6", "21.2", "42.5", "63.7", "85.0", "106", "127", "149", "170"]
    lw_ = 36.0
    cw = (IW - lw_) / len(rates)
    tbox(s, xi + lw_, y, 2 * cw, 8.0, "検証（低射出率側）", 16, color=ORANGE_LINE, bold=True, align="center", margin=0.3)
    tbox(s, xi + lw_ + 2 * cw, y, 8 * cw, 8.0, "学習", 16, color=FRAME, bold=True, align="center", margin=0.3)
    tbox(s, xi + lw_ + 10 * cw, y, 2 * cw, 8.0, "検証（高射出率側）", 16, color=ORANGE_LINE, bold=True, align="center", margin=0.3)
    y += 8.0
    tbox(s, xi, y, lw_, 11.0, "射出率\n(cm^{3}/s)", 15, align="center", anchor="middle", margin=0.3, spacing=0.9)
    for i, r in enumerate(rates):
        test = i < 2 or i >= 10
        b = shape(s, MSO_SHAPE.RECTANGLE, xi + lw_ + i * cw, y, cw, 11.0, fill=ORANGE if test else SKY, line="FFFFFF", lw=1.5)
        text_in(b, r, 18, align="center", anchor="middle", margin=0.3)
    y += 12.5
    t = ("[b]学習[/b]：モデルを作るのに使うデータ（1 条件 5 ショット）　[b]検証[/b]：学習に用いず、予測の答え合わせに使うデータ"
         "（1 条件 3 ショット。低射出率側はキャビティ厚み 1.5・2.0 mm のみ）。外れ値・計測不良のデータは除外")
    tbox(s, xi, y, IW, 18.0, t, NOTE)
    y += 19.0

    y = sub(s, xi, y, "温度シフトファクター")
    wsc, hsc = pic(s, "温度シフト概念図.png", xi + IW - 170.0, y, w=170.0)
    t = ("　Cross-WLF モデルなどの粘度式と同様に、温度の変化によって粘度曲線がずれる（シフトする）性質を表すため、"
         "[r]金型内樹脂最高温度に基づく温度シフトファクター[/r]をモデルに組み込んだ。\n"
         "・PP・粉砕材混合 PP：アレニウス型\n"
         "・GPPS：WLF 型（ガラス転移温度は 100 ℃ と仮定）\n"
         "　粘度方向とせん断速度方向のずれの量は、学習データに最もよく合うように決めた。")
    tbox(s, xi, y, IW - 174.0, hsc, t, 20, hang=1.0, spacing=1.0)
    y = max(y + hsc + 2.0, min_bottom)
    frame(s, x, y0, CW, y)
    return y


def results_band(s, y, title, figs, right, note=None):
    """全幅の帯：左に 3 材料の図（note があれば、その下に注）、右に説明（right(s, x, y, w) が書く）"""
    x, w = M, PAGE_W - 2 * M
    y0 = y
    y = header(s, x, y, w, title) + 2.5
    xs = x + PAD
    labels = ["(1) PP", "(2) 粉砕材混合 PP", "(3) GPPS"]
    hmax, hl = 0.0, 9.0
    for i, (f, lab) in enumerate(zip(figs, labels)):
        xf = xs + i * 131.0
        tbox(s, xf, y, 129.0, hl, lab, 20, bold=True, align="center", anchor="middle", margin=0.2)
        wf, hf = pic(s, f, xf, y + hl, w=129.0)
        hmax = max(hmax, hf + hl)
    if note:
        tbox(s, xs, y + hmax + 0.5, 3 * 131.0 - 2.0, 9.0, note, NOTE)
        hmax += 10.0
    xr = xs + 3 * 131.0 + 4.0
    yr = right(s, xr, y, x + w - PAD - xr)
    y = max(y + hmax, yr) + 2.5
    frame(s, x, y0, w, y)
    return y


def band_learning(s, x, y, w):
    t = ("[b]点[/b]：学習データ（成形条件ごとの平均）、[b]線[/b]：加熱シリンダ温度ごとの予測平均値、[b]帯[/b]：95% 予測区間。"
         "曲線は ΔP_{2} を学習データの中央値、金型内樹脂最高温度を加熱シリンダ温度ごとの中央値に固定して求めたため、"
         "ΔP_{2} が中央値と大きく異なる薄肉条件の点は、曲線からずれる。")
    tbox(s, x, y, w, 30.0, t, NOTE)
    y += 31.0
    bw = w - 30.0
    box(s, x, y, bw, 19.0, "データが密な範囲では予測区間が狭く、データのない低・高せん断速度側では広い", 20, anchor="middle")
    arrow(s, MSO_SHAPE.DOWN_ARROW, x + bw + 8.0, y + 1.0, 14.0, 17.0)
    y += 20.5
    box(s, x, y, w, 12.0, "[r]データの量に応じた予測のばらつきが表現されており、本学習結果は妥当[/r]", 20, anchor="middle", align="center")
    y += 15.0
    we, he = pic(s, "温度シフトの効果.png", x, y, w=120.0)
    t = ("[b]温度シフトファクターの効果[/b]\n温度シフトファクターで換算すると、温度によるデータのずれ（同じ厚みのデータを 1 本の曲線で表したときのずれ）は"
         "[r]換算前の約 4〜5 割に減少[/r]し、3 材料とも、データへの当てはまりが向上した（偶然の差とはいえない大きさ）。\n"
         "※加熱シリンダ温度に基づくシフトファクターでは、PP・粉砕材混合 PP での向上は偶然の差の程度。金型内樹脂最高温度は、"
         "せん断発熱によって射出率とともに変わるため、せん断発熱の影響も含まれる可能性がある。")
    tbox(s, x + 122.0, y, w - 122.0, he, t, 17, spacing=1.0)
    return y + he


def band_verify(s, x, y, w):
    t = ("学習に用いていない射出率条件（検証データ）の粘度を予測し、実測値と比べた。[b]横軸[/b]：実測値、[b]縦軸[/b]：予測値、"
         "[b]縦線[/b]：95% 予測区間。点が破線（予測値 ＝ 実測値）に近いほど正確。")
    tbox(s, x, y, w, 20.0, t, NOTE)
    y += 21.0
    wb = 148.0
    wt = w - wb - 3.0
    rows = [["", "PP", "粉砕材混合 PP", "GPPS"],
            ["平均的な予測誤差^{※}", "約 13%", "約 5%", "約 9%"],
            ["　低射出率側", "約 24%", "約 6%", "約 13%"],
            ["　高射出率側", "約 6%", "約 4%", "約 7%"],
            ["95% 予測区間に入った割合", "89%", "98%", "92%"],
            ["参考：Cross-WLF モデル", "約 72%", "約 59%", "約 97%"]]
    yt = table(s, x, y, [wt - 3 * 46.0, 46.0, 46.0, 46.0], rows, 17)
    tbox(s, x, yt + 0.5, wt, 7.0, "※粘度の対数の二乗平均平方根誤差（RMSE）を、粘度の % に換算した値", 15)
    pic(s, "予測誤差の比較.png", x + w - wb, y, w=wb)
    y = yt + 8.5
    box(s, x, y, wt, 26.0, "高射出率側の誤差は 3 材料とも約 4〜7% と小さく、[r]検証データの 89〜98% が 95% 予測区間に入った[/r]。"
        "Cross-WLF モデルは、キャビティ厚みによる見かけの粘度の違いを表せないため誤差が大きい", 17, anchor="middle", spacing=1.0)
    return y + 26.0


def summary(s, y):
    x, w = M, PAGE_W - 2 * M
    y0 = y
    y = header(s, x, y, w, "8. まとめと今後の課題") + 2.0
    t = ("・金型内センサの計測データとガウス過程回帰を用いて、PP・粉砕材混合 PP・GPPS について、[r]樹脂粘度の予測値と、予測のばらつきの範囲を"
         "同時に求められる[/r]データ駆動型樹脂粘度モデルを構築した。\n"
         "・金型内樹脂最高温度に基づく温度シフトファクターを組み込み、温度による粘度曲線のずれを表現した（ずれは換算前の約 4〜5 割に減少）。\n"
         "・学習に用いていない射出率条件での予測誤差は約 5〜13%（Cross-WLF モデルは約 59〜97%）で、検証データの 89〜98% が 95% 予測区間に入った。\n"
         "・[b]今後の課題[/b]：計測データのない約 150 1/s より小さい低せん断速度域では予測区間が大きく広がり、粉砕材混合 PP では予測値が低せん断速度側で"
         "低下するなど、物理的に妥当でない曲線の形もみられた。低せん断速度側の計測データの追加や、Cross 式などの粘度式とガウス過程回帰を"
         "組み合わせる方法を検討する。")
    tbox(s, x + PAD, y, w - 2 * PAD, 52.0, t, 21, hang=1.0, spacing=1.0)
    y += 52.0
    refs = ("【参考文献】1) 後藤：成形加工’31, 8, 297 (2019)　2) 村田 他：成形加工’シンポジア 25, 24 (2025)　"
            "3) 持橋, 大羽：ガウス過程と機械学習, 講談社 (2019)　4) 内山：機械学習のための確率過程入門, オーム社 (2025)")
    tbox(s, x + PAD, y, w - 2 * PAD, 8.0, refs, 15)
    y += 9.0
    frame(s, x, y0, w, y)
    return y


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(PAGE_W), Mm(PAGE_H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = title_band(s)
    yl = left_column(s, y)
    yr = right_column(s, y, min_bottom=yl)
    y = max(yl, yr) + 5.0
    y = results_band(s, y, "6. 学習結果（樹脂粘度とせん断速度の関係）", ["学習結果_PP.png", "学習結果_MIX.png", "学習結果_GPPS.png"], band_learning,
                     note="※せん断速度が約 150 1/s より小さい範囲は、学習データのない範囲（外挿）。3 温度の帯が重なった部分は濃く見える") + 5.0
    y = results_band(s, y, "7. 予測精度の検証（学習に用いていない射出率条件）", ["予測検証_PP.png", "予測検証_MIX.png", "予測検証_GPPS.png"],
                     band_verify) + 5.0
    y = summary(s, y)
    print(f"左の列の下端 {yl:.1f} mm、右の列の下端 {yr:.1f} mm、最後の下端 {y:.1f} mm（ページの高さ {PAGE_H:.0f} mm）")
    if y > PAGE_H - 4.0:
        print("★ページからはみ出している")
    prs.core_properties.title = "金型内センサデータを用いたデータ駆動型樹脂粘度モデルの開発Ⅲ"
    prs.save(OUT)
    print("保存:", OUT)


if __name__ == "__main__":
    sys.exit(main())
