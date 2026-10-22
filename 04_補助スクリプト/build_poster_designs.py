# -*- coding: utf-8 -*-
"""オリジナルデザインの A0 ポスター（縦 841 × 1189 mm、PowerPoint）を 3 案作る。

  案A「ひと目でわかる」インフォグラフィック型：要点だけ（研究のながれ 3 ステップ、成果の数字、予測曲線、予測精度、まとめ）
  案B「Q&A で読み解く」解説型：疑問に答える形で、報告書B の内容と、温度シフトの効果・苦手な範囲・用語ミニ辞典まで
  案C「材料別に比べる」データ詳細型：3 材料の結果を表の形（予測曲線・温度シフトの効果・予測検証・数値）に並べ、考察まで

読み手：プラスチック成形の知識はあるが、機械学習は知らない人。数値は、報告書B・前のポスターと同じモデル（B_Tff・条件平均の入力）による。
図は make_poster_design_figs.py（図_オリジナル/）と make_poster_figs.py（図/ の見本の切り出し）、報告書B の図（学習結果報告書作成_261015/図/）。
文字・図形は PowerPoint の図形なので、後から PowerPoint で直せる。

使い方（リポジトリ直下で）:
  uv run python 04_補助スクリプト/make_poster_figs.py          （見本の切り出し。前のポスターで作成済みなら不要）
  uv run python 04_補助スクリプト/make_poster_design_figs.py   （3 案の図）
  uv run python 04_補助スクリプト/build_poster_designs.py [A] [B] [C]   （省くと 3 案とも）
出力: 05_…/学習結果報告書作成_261015/ポスター/オリジナル案A_ひと目でわかる版.pptx など
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__))))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from poster_parts import PT_MM, Poster, text_width  # noqa: E402

REPORT = os.path.join(".", "05_材料ごとの機械学習データと予測検証データ", "学習結果報告書作成_261015")
POSTER = os.path.join(REPORT, "ポスター")
FIGO = os.path.join(POSTER, "図_オリジナル")
FIGS = os.path.join(POSTER, "図")
FIGR = os.path.join(REPORT, "図")
W, H = 841.0, 1189.0
TITLE = "金型内センサデータを用いたデータ駆動型樹脂粘度モデルの開発Ⅲ"
TITLE_EN = "Development of Data-Driven Melt Viscosity Model using Data Measured by In-Mold Sensors Ⅲ"
AUTHORS, AFFIL = "〔著者名を記入〕", "〔所属を記入〕"
MATS = [("PP", "PP"), ("MIX", "粉砕材混合 PP"), ("GPPS", "GPPS")]
T_MID = {"PP": "210 ℃", "MIX": "210 ℃", "GPPS": "220 ℃"}
RATES = ["1.6", "2.7", "5.3", "10.6", "21.2", "42.5", "63.7", "85.0", "106", "127", "149", "170"]
REFS = ("1) 後藤：成形加工’31, 8, 297 (2019)　2) 村田 他：成形加工’シンポジア 25, 24 (2025)　3) 持橋, 大羽：ガウス過程と機械学習, 講談社 (2019)　"
        "4) 内山：機械学習のための確率過程入門, オーム社 (2025)")
NUM = "Segoe UI Black"


def fo(name):
    return os.path.join(FIGO, name)


def fs_(name):
    return os.path.join(FIGS, name)


def rate_strip(P, x, y, w, size, test_fill, train_fill, test_txt, train_txt, label_size=None, cell_h=None):
    """射出率 12 水準の帯（両端 2 水準ずつ＝検証、間の 8 水準＝学習）。返り値：下端 [mm]"""
    ls = label_size or size
    cw = w / len(RATES)
    ch = cell_h or size * PT_MM * 2.0
    lh = ls * PT_MM * 1.4
    P.text(x, y, 2 * cw, "検証", ls, h=lh, color=test_txt, bold=True, align="center")
    P.text(x + 2 * cw, y, 8 * cw, "学習", ls, h=lh, color=train_txt, bold=True, align="center")
    P.text(x + 10 * cw, y, 2 * cw, "検証", ls, h=lh, color=test_txt, bold=True, align="center")
    y += lh + 0.8
    for i, r in enumerate(RATES):
        test = i < 2 or i >= 10
        P.box(x + i * cw, y, cw, r, size, h=ch, fill=test_fill if test else train_fill, line="FFFFFF", lw=1.5, pad=(0.3, 0.3, 0, 0),
              color="FFFFFF" if not test or test_fill not in ("FDEBD3", "FCE4D6") else "000000", align="center", anchor="middle", ls=1.0)
    return y + ch


# ================================================================ 案A：ひと目でわかる（インフォグラフィック）
def design_a():
    NAVY, ORANGE, LIGHT, SUB, INK, PALE = "14213D", "F28F1C", "E8EEF7", "5B6577", "1E1E1E", "FDEBD3"
    FONT = "BIZ UDPゴシック"
    P = Poster(W, H, FONT, INK)
    M = 22.0
    CW = W - 2 * M

    def heading(y, num, title, sub=None):
        P.text(M, y, 30, f"[f={NUM}]{num}[/f]", 34, color=ORANGE, ls=1.0)
        _, h = P.text(M + 27, y + 0.6, CW - 27, title, 33, color=NAVY, bold=True, ls=1.0)
        if sub:
            tw = text_width(title, P.base(33, bold=True))
            P.text(M + 27 + tw + 6, y + 5.0, CW - 33 - tw, sub, 18, color=SUB, ls=1.0)
        P.line(M, y + 15.5, M + CW, y + 15.5, color="C9D1DE", lw=1.2)
        return y + 20.0

    # ---------------- 題目
    HH = 134.0
    P.rect(0, 0, W, HH, fill=NAVY)
    P.rect(0, 0, 9, HH, fill=ORANGE)
    P.pic(fo("A_題目の飾り.png"), W - M - 270, 9, w=270)
    P.text(M + 4, 12, CW, "金型内センサデータを用いた", 62, color="FFFFFF", bold=True, ls=1.0)
    P.text(M + 4, 38.5, CW, "データ駆動型樹脂粘度モデルの開発Ⅲ", 62, color="FFFFFF", bold=True, ls=1.0)
    P.text(M + 4, 69, CW, TITLE_EN, 22, color="AFC0DA", ls=1.0)
    P.text(M + 4, 87, CW, AUTHORS, 30, color="FFFFFF", bold=True, ls=1.0)
    P.text(M + 4, 105, CW, AFFIL, 22, color="D6DEEB", ls=1.0)
    P.rect(0, HH, W, 40, fill=ORANGE)
    P.text(0, HH, W, "金型内の計測データから、樹脂粘度の「[c=14213D]予測値[/c]」と「[c=14213D]ばらつきの幅[/c]」を同時に予測する", 40,
           h=40, color="FFFFFF", bold=True, align="center", anchor="middle", ls=1.0)

    # ---------------- 01 研究のながれ
    y = heading(HH + 52, "01", "研究のながれ", "計測 → 機械学習 → 答え合わせ")
    gap = 30.0
    cw = (CW - 2 * gap) / 3
    pad = 9.0
    iw = cw - 2 * pad
    top = y + 2
    steps = [("1", "測る", "計測金型で、流動中の樹脂の[b]圧力と温度[/b]を測り、[b]粘度とせん断速度[/b]を求める"),
             ("2", "学ぶ", "機械学習（[b]ガウス過程回帰[/b]）で、計測データと粘度の関係を学ぶ"),
             ("3", "確かめる", "[b]学習に使っていない射出率[/b]の条件で、予測を答え合わせする")]
    bottoms = []
    for i, (n, verb, desc) in enumerate(steps):
        x = M + i * (cw + gap)
        yy = top + pad
        P.circle(x + pad + 13, yy + 13, 26, n, 36, fill=NAVY, font=NUM, bold=False)
        P.text(x + pad + 32, yy + 2.5, iw - 32, verb, 42, color=NAVY, bold=True, ls=1.0)
        yy += 31
        _, h = P.text(x + pad, yy, iw, desc, 21, ls=1.3)
        yy += max(h, 2 * 21 * PT_MM * 1.3) + 4
        if i == 0:
            w_, h_ = P.pic(fs_("見本_計測金型.png"), x + pad, yy, w=iw)
            yy += h_ + 4
            chips = ["PP・粉砕材混合 PP・GPPS", "厚み 0.5〜2.0 mm", "加熱シリンダ温度 3 水準", "射出率 12 水準"]
            cx = x + pad
            for c in chips:
                tw = P.pill_w(c, 17)
                if cx + tw > x + pad + iw:
                    cx, yy = x + pad, yy + 11
                P.pill(cx, yy, c, 17, fill="FFFFFF", color=NAVY, h=9)
                cx += tw + 3
            yy += 9
        elif i == 1:
            w_, h_ = P.pic(fo("A_GPR概念図.png"), x + pad, yy, w=iw)
            yy += h_ + 3
            _, h = P.box(x + pad, yy, iw, "[c=14213D][b]入力[/b][/c]：せん断速度・樹脂圧力差 ΔP_{2}・金型内樹脂最高温度・加熱シリンダ温度\n"
                         "[c=14213D][b]出力[/b][/c]：粘度の予測値と、[b]95% 予測区間[/b]（ばらつきの幅）", 17, fill="FFFFFF", radius=3,
                         pad=(3, 3, 2.2, 2.2), ls=1.25)
            yy += h
        else:
            yy = rate_strip(P, x + pad, yy + 2, iw, 15, test_fill=ORANGE, train_fill=NAVY, test_txt="C46D00", train_txt=NAVY, label_size=17,
                            cell_h=14) + 1
            P.text(x + pad, yy, iw, "射出率（cm^{3}/s）", 15, color=SUB, align="center", ls=1.0)
            yy += 9
            t = ("[c=14213D]■[/c] [b]学習[/b]：間の 8 水準（1 条件 5 ショット）\n"
                 "[c=F28F1C]■[/c] [b]検証[/b]：両端の 4 水準（1 条件 3 ショット。低射出率側は厚み 1.5・2.0 mm のみ）")
            _, h = P.text(x + pad, yy, iw, t, 18.5, ls=1.35, after=2)
            yy += h + 3
            w_, h_ = P.pic(fo("A_検証の位置_PP.png"), x + pad, yy, w=iw)
            yy += h_ + 1
            _, h = P.text(x + pad, yy, iw, "検証の条件は、学習データの範囲の[b]外側[/b]（PP、厚み 2.0 mm の例。点は成形条件ごとの平均）", 16.5, color=SUB,
                          ls=1.3)
            yy += h
        bottoms.append(yy + pad)
    card_b = max(bottoms)
    for i in range(3):
        x = M + i * (cw + gap)
        card = P.rect(x, top, cw, card_b - top, fill=LIGHT, radius=6)
        P.s.shapes._spTree.remove(card._element)
        P.s.shapes._spTree.insert(2 + i, card._element)                 # カードを、中身より後ろに回す
        if i < 2:
            P.shape(MSO_SHAPE.CHEVRON, x + cw + gap / 2 - 8, top + (card_b - top) / 2 - 15, 16, 30, fill=ORANGE)

    # ---------------- 数字で見る成果（帯）
    y = card_b + 10
    band_h = 100.0
    P.rect(0, y, W, band_h, fill=NAVY)
    P.text(M, y + 6, CW, "数字で見る成果", 22, color="AFC0DA", bold=True, align="center", ls=1.0)
    kpis = [("[f=Segoe UI Black]3[/f][s=36][c=FFFFFF] 材料[/c][/s]", "PP・粉砕材混合 PP・GPPS で\nモデルを構築"),
            ("[f=Segoe UI Black]5[/f][s=56]〜[/s][f=Segoe UI Black]13[/f][s=40][c=FFFFFF]%[/c][/s]", "学習に使っていない条件での\n平均的な予測誤差"),
            ("[f=Segoe UI Black]89[/f][s=56]〜[/s][f=Segoe UI Black]98[/f][s=40][c=FFFFFF]%[/c][/s]", "実測値が 95% 予測区間に\n入った割合"),
            ("[f=Segoe UI Black]1/5[/f][s=36][c=FFFFFF] 以下[/c][/s]", "予測誤差は、従来の\nCross-WLF モデルの 1/5 以下")]
    tw = CW / 4
    for i, (num, lab) in enumerate(kpis):
        x = M + i * tw
        P.text(x, y + 16, tw, num, 80, color=ORANGE, bold=True, align="center", ls=1.0)
        P.text(x, y + 52, tw, lab, 21, color="FFFFFF", align="center", ls=1.3)
        if i:
            P.line(x, y + 22, x, y + band_h - 12, color="3A4A6B", lw=1.5)
    y += band_h + 10

    # ---------------- 02 成果① 予測曲線
    y = heading(y, "02", "成果①　厚みごとの粘度と、そのばらつきの幅を予測できた")
    fw = 252.0
    fg = (CW - 3 * fw) / 2
    for i, (key, name) in enumerate(MATS):
        x = M + i * (fw + fg)
        _, pw, ph = P.pill(x + 2, y, name, 24, fill=NAVY, h=12)
        P.text(x + pw + 6, y + 2.5, fw - pw - 6, f"加熱シリンダ温度 {T_MID[key]}", 17, color=SUB, ls=1.0)
        P.pic(fo(f"A_予測曲線_{key}.png"), x, y + 14, w=fw)
    y += 14 + 236 + 1
    lw_, lh_ = P.pic(fo("A_凡例_予測曲線.png"), M + (CW - 600) / 2, y, w=600)
    y += lh_ + 1
    t = ("線と帯は、厚みごとの ΔP_{2}・金型内樹脂最高温度（成形条件ごとの平均）を入力して求めた予測。"
         "[b]白抜きの点（学習に使っていない条件）も、曲線の近くに乗っている[/b]。データのない範囲では、帯（ばらつきの幅）が広がる")
    _, h = P.text(M + 20, y, CW - 40, t, 18.5, color=INK, align="center", ls=1.3)
    y += h + 8

    # ---------------- 03 成果② 予測精度
    y = heading(y, "03", "成果②　学習に使っていない条件でも、よく当たった")
    pw_, pg = 168.0, 10.0
    for i, (key, name) in enumerate(MATS):
        x = M + i * (pw_ + pg)
        P.text(x, y, pw_, name, 21, color=NAVY, bold=True, align="center", ls=1.0)
        P.pic(fo(f"A_予測検証_{key}.png"), x, y + 9, w=pw_)
    rx = M + 3 * (pw_ + pg) + 6
    rw = M + CW - rx
    P.text(rx, y, rw, "従来の粘度式との比較", 21, color=NAVY, bold=True, align="center", ls=1.0)
    _, bh = P.pic(fo("A_誤差の比較.png"), rx, y + 9, w=rw)
    y += 9 + 166 + 2
    t = ("横軸が実測値、縦軸が予測値。点が破線に近いほど正確で、縦線が 95% 予測区間（ばらつきの幅）。"
         "Cross-WLF モデルは、キャビティ厚みによる見かけの粘度の違いを表せないため、誤差が大きいと考えられる")
    _, h = P.text(M, y, CW, t, 18.5, color=INK, ls=1.3)
    y += h + 8

    # ---------------- まとめ
    bw = (CW - 12) / 2
    t1 = ("[s=26][c=14213D][b]わかったこと[/b][/c][/s]\n"
          "・金型内センサの計測データとガウス過程回帰で、樹脂粘度の[b]予測値[/b]と[b]ばらつきの幅（95% 予測区間）[/b]を同時に求めるモデルを、3 材料で構築\n"
          "・学習に使っていない射出率条件で、予測誤差は約 5〜13%、実測値の 89〜98% が 95% 予測区間に入った\n"
          "・金型内樹脂最高温度に基づく温度シフトファクターで、温度による粘度曲線のずれも表現")
    t2 = ("[s=26][c=C46D00][b]今後の課題[/b][/c][/s]\n"
          "・学習データのない低せん断速度側（約 150 1/s より小さい範囲）では、ばらつきの幅が大きく広がり、曲線の形が不自然になることもある\n"
          "・低せん断速度側の計測データの追加や、Cross 式などの粘度式とガウス過程回帰の組み合わせを検討する")
    h1 = P.text_height(t1, P.base(19), bw - 12, 1.3, 3, 1.0)
    h2 = P.text_height(t2, P.base(19), bw - 12, 1.3, 3, 1.0)
    hh = max(h1, h2) + 9
    P.box(M, y, bw, t1, 19, h=hh, fill=LIGHT, radius=5, pad=(6, 6, 4.5, 4.5), ls=1.3, after=3, hang=1.0)
    P.box(M + bw + 12, y, bw, t2, 19, h=hh, fill=PALE, radius=5, pad=(6, 6, 4.5, 4.5), ls=1.3, after=3, hang=1.0)
    y += hh + 5
    P.text(M, y, CW, "【参考文献】" + REFS, 14, color=SUB, ls=1.2)
    return P, y + 8


def chips(P, x, y, w, items, size, fill, color, h=None, gap=3.0):
    """札（pill）を横に並べ、入りきらなければ折り返す。返り値：高さ"""
    h = h or size * PT_MM * 1.55
    cx, cy = x, y
    for c in items:
        tw = P.pill_w(c, size)
        if cx + tw > x + w and cx > x:
            cx, cy = x, cy + h + 2
        P.pill(cx, cy, c, size, fill=fill, color=color, h=h)
        cx += tw + gap
    return cy + h - y


# ================================================================ 案B：Q&A で読み解く（解説型）
def design_b():
    TD, TM, TL, CORAL, SAND, INK, BG, SUB = "004D40", "00796B", "E0F2F1", "F4511E", "FFF3E0", "263238", "EEF3F2", "546E7A"
    FONT, HFONT = "游ゴシック Medium", "游ゴシック"
    P = Poster(W, H, FONT, INK)
    P.rect(0, 0, W, H, fill=BG)
    M, G = 18.0, 12.0
    CW = W - 2 * M
    col_w = (CW - G) / 2
    PAD = 9.0
    SH = (3.0, 0.8, 0.12)

    def card(x, y, w, num, q, a, body):
        """Q の札。body(x, y, w) は中身を書いて高さを返す。返り値：(下端, 背景の図形)"""
        mk = P.mark()
        cx, cy = x + PAD, y + PAD
        P.circle(cx + 11, cy + 11, 22, f"Q{num}", 17, fill=TM, font=NUM, bold=False)
        _, qh = P.text(cx + 27, cy + 2.5, w - 2 * PAD - 27, q, 27, color=TD, bold=True, font=HFONT, ls=1.15)
        yy = cy + max(22, qh + 4) + 2
        P.text(cx + 1, yy - 0.3, 14, "A.", 21, color=CORAL, font=NUM, ls=1.25)
        _, ah = P.text(cx + 14, yy, w - 2 * PAD - 14, a, 21, color=INK, bold=True, font=HFONT, ls=1.3)
        yy += ah + 5
        bh = body(cx, yy, w - 2 * PAD)
        bottom = yy + bh + PAD
        bg = P.rect(x, y, w, bottom - y, fill="FFFFFF", radius=5, shadow=SH)
        P.to_back(bg, mk)
        return bottom, bg

    def row(y, specs):
        """カードを横に並べ、高さをそろえる。specs：[(x, w, num, q, a, body), ...]。返り値：下端"""
        res = [card(x, y, w, num, q, a, body) for x, w, num, q, a, body in specs]
        bottom = max(b for b, _ in res)
        for _, bg in res:
            bg.height = int((bottom - y) * 36000)
        return bottom

    # ---------------- 題目
    HH = 100.0
    P.rect(0, 0, W, HH, fill="FFFFFF")
    P.rect(M, 13, 7, 72, fill=CORAL)
    P.text(M + 14, 10, 520, "金型内センサデータを用いた", 50, color=TD, bold=True, font=HFONT, ls=1.0)
    P.text(M + 14, 31, 520, "データ駆動型樹脂粘度モデルの開発Ⅲ", 50, color=TD, bold=True, font=HFONT, ls=1.0)
    P.text(M + 14, 55, 540, TITLE_EN, 17, color=SUB, ls=1.0)
    P.text(M + 14, 66, 520, AUTHORS, 25, color=INK, bold=True, font=HFONT, ls=1.0)
    P.text(M + 14, 79.5, 520, AFFIL, 18, color=SUB, ls=1.0)
    gx = M + 548
    guide = ("[s=22][b][c=F4511E]このポスターの読み方[/c][/b][/s]\n"
             "・[b]Q1 → Q7[/b] の順に、研究の流れをたどれます\n"
             "・[c=F4511E][b]A.[/b][/c] の太字だけ読めば、要点がわかります\n"
             "・専門用語は、右下の「[b]用語ミニ辞典[/b]」へ")
    P.box(gx, 11, M + CW - gx, guide, 17, h=76, fill=SAND, line="F6B48F", lw=1.2, radius=5, pad=(6, 5, 4, 3), ls=1.35, hang=1.0,
          anchor="middle")
    P.rect(0, HH, W, 2.2, fill=TM)
    RG = 7.0                                                         # 段の間

    # ---------------- 1 段目：Q1・Q2
    xl, xr = M, M + col_w + G

    def q1(x, y, w):
        t = ("射出成形 CAE の樹脂粘度モデル^{1)}（Cross-WLF モデルなど）は、温度とせん断速度が決まると粘度が 1 つの値に決まる式で、樹脂の製造ロット間の分子量分布や、"
             "成形工程の熱履歴による[b]ばらつきは考慮されていない[/b]。本研究では、金型内の計測データと機械学習を用いて、粘度の"
             "[b]予測値[/b]と[b]ばらつきの幅[/b]を同時に求めるモデルの構築を目指す。")
        _, h = P.text(x, y, w, t, 18, ls=1.42)
        yy = y + h + 3
        _, hi = P.pic(fo("B_従来と本研究.png"), x, yy, w=w)
        yy += hi + 4
        t = ("[b]本報のポイント[/b]　① 3 材料（PP・粉砕材混合 PP・GPPS）で構築　② 温度シフトファクターを組み込み　"
             "③ 学習に使っていない射出率の条件で、予測精度を検証")
        _, h = P.box(x, yy, w, t, 16.5, fill=TL, radius=3, pad=(4, 4, 2.5, 2.5), ls=1.4)
        return yy + h - y

    def q2(x, y, w):
        w1 = 208.0
        _, h1 = P.pic(fs_("見本_キャビティの模式図.png"), x, y, w=w1)
        rw = w - w1 - 6
        _, h2 = P.pic(fs_("算出式.png"), x + w1 + 6, y + 4, w=rw)
        yy = y + max(h1, h2 + 4) + 3
        t = ("・センサ①と②の圧力が立ち上がる[b]時間の差 Δt[/b] → 樹脂の流量 Q\n"
             "・②が立ち上がった時刻の①と②の[b]圧力差 ΔP_{1}[/b] → 粘度 η（ニュートン流体の計算式^{2)}）\n"
             "・L：センサ間距離（17.5 mm）、B：流路幅、H：キャビティ厚み")
        _, h = P.text(x, yy, w, t, 16.5, ls=1.38, hang=1.0)
        yy += h + 3
        hh = chips(P, x, yy, w, ["PP・粉砕材混合 PP・GPPS", "厚み 0.5〜2.0 mm", "加熱シリンダ温度 3 水準", "射出率 1.6〜170 cm^{3}/s（12 水準）"], 15,
                   fill=TL, color=TD)
        return yy + hh - y

    y = row(HH + 2.2 + 8, [(xl, col_w, 1, "なぜ、新しい粘度モデルが必要？", "実際の成形では、粘度が材料のロットや熱履歴でばらつくから", q1),
                  (xr, col_w, 2, "粘度は、どうやって測った？", "2 つの圧力センサの間を樹脂が流れる時間と、圧力の差から求めた", q2)])

    # ---------------- 2 段目：Q3・Q4
    def q3(x, y, w):
        wg = 214.0
        _, hg = P.pic(fo("B_GPR概念図.png"), x, y, w=wg)
        t = ("・式の形（Cross 式など）を決めずに、データから関係を学ぶ手法^{3)}\n"
             "・[b]データの多い所は幅が狭く、少ない所は幅が広い[/b]\n"
             "・幅は [b]95% 予測区間[/b]：同じ条件で測り直したとき、計測値の 95% が入ると予測される範囲")
        P.text(x + wg + 6, y + 3, w - wg - 6, t, 16.5, ls=1.38, hang=1.0, after=4)
        yy = y + hg + 4
        bh, aw, mw = 31.0, 11.0, 88.0
        iw_ = (w - mw - 2 * aw - 8) / 2
        P.box(x, yy, iw_, "[b]入力[/b]：せん断速度 γ̇、樹脂圧力差 ΔP_{2}、金型内樹脂最高温度、加熱シリンダ温度", 15, h=bh, fill=TL, radius=3,
              pad=(3, 3, 1, 1), anchor="middle", ls=1.25)
        P.shape(MSO_SHAPE.RIGHT_ARROW, x + iw_ + 2, yy + bh / 2 - 6, aw, 12, fill=CORAL)
        P.box(x + iw_ + aw + 4, yy, mw, "ガウス過程回帰\n（温度シフト付き）", 16, h=bh, fill=TM, color="FFFFFF", bold=True, font=HFONT, radius=3,
              pad=(1, 1, 1, 1), align="center", anchor="middle", ls=1.2)
        P.shape(MSO_SHAPE.RIGHT_ARROW, x + iw_ + aw + mw + 6, yy + bh / 2 - 6, aw, 12, fill=CORAL)
        P.box(x + w - iw_, yy, iw_, "[b]出力[/b]：粘度の[b]予測値[/b]と、[b]95% 予測区間[/b]（ばらつきの幅）", 15, h=bh, fill=TL, radius=3,
              pad=(3, 3, 1, 1), anchor="middle", ls=1.25)
        yy += bh + 3
        t = "※η・γ̇・ΔP_{2} は対数で扱った。γ̇・ΔP_{2}・金型内樹脂最高温度は、同じ成形条件のショットの平均値を使った"
        _, h = P.text(x, yy, w, t, 14.5, color=SUB, ls=1.3)
        return yy + h - y

    def q4(x, y, w):
        wc = 122.0
        _, hc = P.pic(fo("B_温度シフト概念図.png"), x, y, w=wc)
        t = ("温度が変わると、粘度曲線は横と縦にずれる。このずれを、Cross-WLF モデルと同じ考え方の[b]温度シフトファクター[/b]で表した。\n"
             "・PP・粉砕材混合 PP：アレニウス型\n・GPPS：WLF 型（ガラス転移温度 100 ℃ と仮定）\n"
             "・温度には[b]金型内樹脂最高温度[/b]を使用。ずれの量は、データに最もよく合うように決めた")
        P.text(x + wc + 6, y + 1, w - wc - 6, t, 16.5, ls=1.38, hang=1.0, after=2)
        yy = y + hc + 3
        _, hr = P.pic(fo("B_温度のずれ_PP.png"), x, yy, w=w)
        yy += hr + 1
        t = ("PP の例（点＝成形条件、横線＝温度ごとの平均）。換算前は、低温ほど上・高温ほど下にずれていたが、換算後はほぼ 0。"
             "温度によるずれ：PP 6.9 → 3.0%、粉砕材混合 PP 7.3 → 3.7%、GPPS 12.8 → 5.0%")
        _, h = P.text(x, yy, w, t, 14.5, color=SUB, ls=1.3)
        return yy + h - y

    y = row(y + RG, [(xl, col_w, 3, "機械学習（ガウス過程回帰）って何？", "データから入出力の関係を学ぶ方法。予測値と一緒に「ばらつきの幅」も出せる", q3),
                     (xr, col_w, 4, "温度の影響は、どう表した？", "「温度シフトファクター」で曲線をずらし、温度によるずれを約半分（換算前の約 4〜5 割）に減らした", q4)])

    # ---------------- 3 段目：Q5（全幅）
    def q5(x, y, w):
        fw = 253.0
        fg = (w - 3 * fw) / 2
        for i, (key, name) in enumerate(MATS):
            xx = x + i * (fw + fg)
            _, pw, _ = P.pill(xx + 2, y, name, 18, fill=TM, h=9.5)
            P.text(xx + pw + 6, y + 1.8, fw - pw - 6, f"加熱シリンダ温度 {T_MID[key]}", 15, color=SUB, ls=1.0)
            P.pic(fo(f"B_予測曲線_{key}.png"), xx, y + 11, w=fw)
        yy = y + 11 + 172 + 1
        _, lh = P.pic(fo("B_凡例_予測曲線.png"), x + (w - 600) / 2, yy, w=600)
        yy += lh + 1.5
        t = ("線と帯は、各厚みの ΔP_{2}・金型内樹脂最高温度（成形条件ごとの平均）を入力した予測。同じせん断速度でも薄肉ほど見かけの粘度が高く、"
             "その違いも表せている。[b]白抜きの点（学習に使っていない条件）も、曲線の近くに乗っている[/b]")
        _, h = P.text(x, yy, w, t, 16.5, ls=1.38)
        return yy + h - y

    y = row(y + RG, [(M, CW, 5, "どんな予測ができる？", "厚みごとの粘度曲線と、そのばらつきの幅を、3 材料で予測できた", q5)])

    # ---------------- 4 段目：Q6（全幅）
    def q6(x, y, w):
        lw_ = 300.0
        yy = rate_strip(P, x, y, lw_, 13.5, test_fill=CORAL, train_fill=TM, test_txt=CORAL, train_txt=TD, label_size=15, cell_h=11)
        _, h = P.text(x, yy + 0.8, lw_, "射出率（cm^{3}/s）。両端の 4 水準は学習に使わず、答え合わせに使った", 13.5, color=SUB, ls=1.25)
        yy += h + 3
        rows = [["", "PP", "粉砕材混合 PP", "GPPS"],
                ["平均的な予測誤差", "約 13%", "約 5%", "約 9%"],
                ["　低射出率側", "約 24%", "約 6%", "約 13%"],
                ["　高射出率側", "約 6%", "約 4%", "約 7%"],
                ["予測区間に入った割合", "89%", "98%", "93%"],
                ["参考：Cross-WLF の誤差", "約 72%", "約 59%", "約 97%"]]
        yb = P.table(x, yy, [111, 63, 63, 63], rows, 15, head_fill=TD, fills=("FFFFFF", "F1F6F5"), line="D5DEDC", lw=0.75)
        _, h = P.text(x, yb + 1, lw_, "誤差は、粘度の対数の二乗平均平方根誤差（RMSE）を % に換算した値", 12.5, color=SUB, ls=1.2)
        t = ("[b][c=F4511E]読み方[/c][/b]　誤差は小さいほど正確。予測区間に入った割合は [b]95% に近いほど[/b]、ばらつきの幅の見積もりが正しい。"
             "右の図は、点が破線に近いほど正確で、縦線が 95% 予測区間")
        _, h2 = P.box(x, yb + 1 + h + 3, lw_, t, 15, fill=SAND, radius=3, pad=(4, 4, 2.5, 2.5), ls=1.38)
        h += 3 + h2
        pw = 130.0
        pg = (w - lw_ - 3 * pw) / 3
        for i, (key, name) in enumerate(MATS):
            xx = x + lw_ + pg + i * (pw + pg)
            P.text(xx, y, pw, name, 17, color=TD, bold=True, font=HFONT, align="center", ls=1.0)
            P.pic(fo(f"B_予測検証_{key}.png"), xx, y + 7, w=pw)
        return max(yb + 1 + h, y + 7 + 127) - y

    y = row(y + RG, [(M, CW, 6, "どのくらい当たる？",
                      "学習に使っていない条件で、誤差は約 5〜13%。実測値の 89〜98% が 95% 予測区間に入った（Cross-WLF モデルの誤差は約 59〜97%）", q6)])

    # ---------------- 5 段目：Q7・まとめ・用語
    w7 = 352.0

    def q7(x, y, w):
        wf = 165.0
        _, hf = P.pic(fo("B_苦手な範囲.png"), x, y, w=wf)
        t = ("粉砕材混合 PP（210 ℃、全厚み共通の曲線）の例。灰色は学習データのない範囲。約 150 1/s より小さい範囲では、"
             "予測区間が大きく広がり、粘度が下がる不自然な曲線の形も見られた。\n"
             "→ 低せん断速度側の計測データの追加や、Cross 式などの粘度式との組み合わせを検討する")
        _, h = P.text(x + wf + 6, y + 1, w - wf - 6, t, 15.5, ls=1.38, after=4)
        return max(hf, h) + 0.0

    yb7, bg7 = card(M, y + RG, w7, 7, "苦手なことは？", "学習データのない低せん断速度側では、ばらつきの幅が大きく広がる", q7)
    sx = M + w7 + G
    sw = (M + CW - sx - G) * 0.54
    gx = sx + sw + G
    gw = M + CW - gx
    t = ("[s=24][b][c=004D40]まとめ[/c][/b][/s]\n"
         "・金型内センサの計測データとガウス過程回帰で、樹脂粘度の[b]予測値[/b]と[b]ばらつきの幅[/b]を同時に求めるモデルを、3 材料で構築した\n"
         "・学習に使っていない射出率の条件で、予測誤差は約 5〜13%、実測値の 89〜98% が 95% 予測区間に入った\n"
         "・温度シフトファクターで、温度によるずれを換算前の約 4〜5 割に減らした")
    s_s, hs = P.box(sx, y + RG, sw, t, 18.5, fill="FFFFFF", radius=5, pad=(7, 7, 6, 6), ls=1.4, hang=1.0, after=4, shadow=SH)
    t = ("[s=22][b][c=F4511E]用語ミニ辞典[/c][/b][/s]\n"
         "[b]機械学習[/b]：データから、入力と出力の関係を自動で学ぶ方法\n"
         "[b]学習データ・検証データ[/b]：モデルを作るデータと、答え合わせに使うデータ\n"
         "[b]95% 予測区間[/b]：測り直したとき、計測値の 95% が入ると予測される範囲\n"
         "[b]外挿[/b]：学習データの範囲の外を予測すること。予測の幅が広がりやすい\n"
         "[b]温度シフトファクター[/b]：温度による粘度曲線のずれ（横・縦）を表す係数")
    s_g, hg = P.box(gx, y + RG, gw, t, 16, fill=SAND, radius=5, pad=(6, 6, 6, 6), ls=1.38, after=4)
    bottom = max(yb7, y + RG + hs, y + RG + hg)
    for s_ in (bg7, s_s, s_g):
        s_.height = int((bottom - y - RG) * 36000)
    y = bottom + 5
    _, h = P.text(M, y, CW, "【参考文献】" + REFS, 13, color=SUB, ls=1.2)
    return P, y + h


# ================================================================ 案C：材料別に比べる（データ詳細型）
def design_c():
    CH, ST, LT, AC, INK, SUB, LINE = "22303C", "4A6274", "EEF1F4", "D1495B", "1F2933", "52606D", "CBD2D9"
    MCOL = {"PP": "3D7EBF", "MIX": "2E8B57", "GPPS": "8E5BB0"}
    FONT = "Meiryo UI"
    P = Poster(W, H, FONT, INK)
    M = 18.0
    CW = W - 2 * M
    NUMF = "Segoe UI"

    def heading(x, y, w, num, title):
        """節の見出し：番号の四角と題、下に細い線。返り値：中身を書き始める y"""
        P.box(x, y, 13, num, 22, h=13, fill=CH, color="FFFFFF", bold=True, font=NUMF, pad=(0, 0, 0, 0), align="center", anchor="middle", ls=1.0)
        P.text(x + 16, y + 0.8, w - 16, title, 25, color=CH, bold=True, ls=1.0)
        P.line(x, y + 15.5, x + w, y + 15.5, color=CH, lw=1.5)
        return y + 19.5

    # ---------------- 題目
    HH = 118.0
    P.rect(0, 0, W, HH, fill=CH)
    P.text(M, 11, CW, TITLE, 52, color="FFFFFF", bold=True, ls=1.0)
    P.text(M, 33, CW, TITLE_EN, 19, color="B8C2CC", ls=1.0)
    P.text(M, 45, CW / 2, AUTHORS, 25, color="FFFFFF", bold=True, ls=1.0)
    P.text(M, 57.5, CW / 2, AFFIL, 18, color="B8C2CC", ls=1.0)
    abst = ("[b][c=EDAE49]要旨[/c][/b]　金型内センサで計測した樹脂圧力・温度のデータと機械学習（ガウス過程回帰）を用いて、樹脂粘度の[b]予測値[/b]と"
            "[b]95% 予測区間（ばらつきの幅）[/b]を同時に求めるモデルを、PP・粉砕材混合 PP・GPPS で構築した。\n学習に使っていない射出率の条件で、"
            "予測誤差は約 5〜13%、実測値の 89〜98% が 95% 予測区間に入った。")
    P.box(M, 71, CW, abst, 18.5, h=40, fill="33465A", color="FFFFFF", radius=3, pad=(6, 6, 2, 2), anchor="middle", ls=1.38)

    # ---------------- 1 段目：背景・計測・モデル
    y0 = HH + 9
    cg = 11.0
    cw = (CW - 2 * cg) / 3
    xs = [M + i * (cw + cg) for i in range(3)]
    # 1 背景と目的
    x = xs[0]
    y = heading(x, y0, cw, "1", "背景と目的")
    t = ("射出成形 CAE で用いられる樹脂粘度モデル^{1)} には、樹脂の製造ロット間の分子量分布のばらつきや、成形工程の熱履歴による"
         "ばらつきが考慮されていない。本研究では、計測金型と機械学習を用いて、樹脂粘度の[b]予測値[/b]に加えて、[b]予測のばらつきの範囲[/b]も"
         "同時に求められる、データ駆動型樹脂粘度モデルの構築を目的としている。")
    _, h = P.text(x, y, cw, t, 18, ls=1.42)
    y += h + 4
    t = ("[b]本報の内容[/b]\n・3 材料（PP・粉砕材混合 PP・GPPS）でモデルを構築\n・Cross-WLF モデルと同様の[b]温度シフトファクター[/b]を組み込み\n"
         "・[b]学習に使っていない射出率[/b]の条件で予測精度を検証し、Cross-WLF モデルと比較")
    _, h = P.box(x, y, cw, t, 17, fill=LT, radius=3, pad=(5, 5, 3.5, 3.5), ls=1.38, hang=1.0)
    y += h + 4
    _, h4 = P.pic(os.path.join(FIGR, "Fig4.png"), x, y, w=cw)
    y += h4 + 1
    _, h = P.text(x, y, cw, "計測した粘度（報告書の図。左から PP・粉砕材混合 PP・GPPS）。せん断速度の増加・温度の上昇で粘度が下がり、"
                  "高せん断速度側はキャピラリーレオメータとおおむね一致", 15, color=SUB, ls=1.32)
    b1 = y + h
    # 2 計測方法
    x = xs[1]
    y = heading(x, y0, cw, "2", "計測方法と成形条件")
    w1 = cw * 0.49
    _, h1 = P.pic(fs_("見本_計測金型.png"), x, y, w=w1)
    _, h2 = P.pic(fs_("見本_キャビティ入れ子.png"), x + w1 + 3, y, w=cw - w1 - 3)
    y += max(h1, h2) + 3
    w3 = cw * 0.56
    _, h3 = P.pic(fs_("見本_キャビティの模式図.png"), x, y, w=w3)
    _, h5 = P.pic(fs_("算出式.png"), x + w3 + 3, y + 2, w=cw - w3 - 3)
    y += max(h3, h5 + 2) + 2
    t = ("・Δt：①と②の圧力の立ち上がり時間の差 → 流量 Q　・ΔP_{1}：②の立ち上がり時の①と②の圧力差 → 粘度 η^{2)}\n"
         "・ΔP_{2}：②の最大値時の①と②の圧力差、金型内樹脂最高温度：②に樹脂が届いた時の最高温度 → 機械学習の入力")
    _, h = P.text(x, y, cw, t, 15.5, ls=1.36, hang=1.0)
    y += h + 3
    rows = [["", "PP", "粉砕材混合 PP", "GPPS"],
            ["加熱シリンダ温度", "190・210・230 ℃", "190・210・230 ℃", "200・220・240 ℃"],
            ["キャビティ厚み", "0.5〜2.0 mm（4 水準）", "0.5〜2.0 mm（4 水準）", "1.0〜2.0 mm（3 水準）"],
            ["射出率", "1.6〜170 cm^{3}/s の 12 水準（3 材料共通）", "", ""]]
    cws = [50.0, (cw - 50) / 3, (cw - 50) / 3, (cw - 50) / 3]
    yb = P.table(x, y, cws, rows[:3], 14.5, head_fill=CH, fills=("FFFFFF", LT), line=LINE, lw=0.75)
    P.box(x, yb, cws[0], rows[3][0], 14.5, h=14.5 * PT_MM * 1.9, fill="FFFFFF", line=LINE, lw=0.75, pad=(1.8, 1.8, 0, 0), bold=True,
          anchor="middle", ls=1.0)
    P.box(x + cws[0], yb, cw - cws[0], rows[3][1], 14.5, h=14.5 * PT_MM * 1.9, fill="FFFFFF", line=LINE, lw=0.75, pad=(1.8, 1.8, 0, 0),
          align="center", anchor="middle", ls=1.0)
    b2 = yb + 14.5 * PT_MM * 1.9
    # 3 機械学習モデル
    x = xs[2]
    y = heading(x, y0, cw, "3", "機械学習モデル（ガウス過程回帰）")
    _, hg = P.pic(fo("C_GPR概念図.png"), x, y, w=cw)
    y += hg + 2
    t = ("・式の形を決めずにデータから関係を学び、[b]予測値と 95% 予測区間[/b]を出す^{3)}。カーネル関数^{4)}は、入力ごとに影響の大きさを調整する ARD カーネルに、"
         "直線的な成分と、ショット間のばらつきを表す成分を加えた\n"
         "・[b]入力[/b]：せん断速度 γ̇、ΔP_{2}、金型内樹脂最高温度、加熱シリンダ温度（γ̇・ΔP_{2}・温度は、成形条件ごとのショットの平均）\n"
         "・[b]温度シフトファクター[/b]：金型内樹脂最高温度で、粘度方向とせん断速度方向にずらす（PP・粉砕材混合 PP：アレニウス型、GPPS：WLF 型）")
    _, h = P.text(x, y, cw, t, 15.5, ls=1.36, hang=1.0, after=2)
    y += h + 3
    y = rate_strip(P, x, y, cw, 12.5, test_fill=AC, train_fill=ST, test_txt=AC, train_txt=CH, label_size=14, cell_h=10.5)
    _, h = P.text(x, y + 0.8, cw, "射出率（cm^{3}/s）。学習：1 条件 5 ショット、検証：1 条件 3 ショット（低射出率側は厚み 1.5・2.0 mm のみ）", 13,
                  color=SUB, ls=1.3)
    b3 = y + 0.8 + h
    y = max(b1, b2, b3) + 14

    # ---------------- 2 段目：結果の表（材料 × 結果）
    y = heading(M, y, CW, "4", "結果：3 材料の比較")
    lab_w, gap = 26.0, 6.0
    fw = (CW - lab_w - 3 * gap) / 3
    fx = [M + lab_w + gap + i * (fw + gap) for i in range(3)]
    tab_h = 19.0
    for i, (key, name) in enumerate(MATS):
        P.rect(fx[i], y, fw, tab_h, fill=MCOL[key], radius=3)
        P.text(fx[i] + 5, y + 2, fw - 10, name, 24, h=tab_h - 4, color="FFFFFF", bold=True, anchor="middle", ls=1.0)
        temps = "200・220・240 ℃" if key == "GPPS" else "190・210・230 ℃"
        thick = "1.0〜2.0 mm" if key == "GPPS" else "0.5〜2.0 mm"
        P.text(fx[i] + 5, y + 2, fw - 10, f"{temps}｜厚み {thick}", 14, h=tab_h - 4, color="FFFFFF", align="right", anchor="middle", ls=1.0)
    y += tab_h + 3

    def row_label(y1, y2, text, sub=None):
        P.rect(M, y1, lab_w, y2 - y1, fill=CH, radius=2)
        s = P.s.shapes.add_textbox(0, 0, 10, 10)
        s.left, s.top = int(M * 36000), int(y1 * 36000)
        s.width, s.height = int(lab_w * 36000), int((y2 - y1) * 36000)
        P.frame(s, text + (f"\n[s=13]{sub}[/s]" if sub else ""), 21, color="FFFFFF", bold=True, align="center", anchor="middle", ls=1.15, vert=True)

    # (a) 予測曲線
    ya = y
    for i, (key, _) in enumerate(MATS):
        P.pic(fo(f"C_予測曲線_{key}.png"), fx[i], y, w=fw)
    y += 214 + 1
    _, lh = P.pic(fo("C_凡例_予測曲線.png"), M + lab_w + gap + (3 * fw + 2 * gap - 600) / 2, y, w=600)
    y += lh + 1
    _, h = P.text(fx[0], y, 3 * fw + 2 * gap, "加熱シリンダ温度 210 ℃（GPPS は 220 ℃）。線と帯は、各厚みの ΔP_{2}・金型内樹脂最高温度（成形条件ごとの平均）を入力した予測。"
                  "白抜きの点（学習に使っていない条件）も曲線の近くに乗り、データのない範囲では帯（ばらつきの幅）が広がる", 15, color=INK, align="center",
                  ls=1.3)
    y += h + 4
    row_label(ya, y - 2, "予測曲線")
    # (b) 温度シフトの効果
    yb_ = y
    for i, (key, _) in enumerate(MATS):
        P.pic(fo(f"C_温度のずれ_{key}.png"), fx[i], y, w=fw)
    y += 104 + 1
    _, h = P.text(fx[0], y, 3 * fw + 2 * gap, "縦軸：同じ厚みの全温度のデータに 1 本の曲線を当てはめたときの、曲線からのずれ。点＝成形条件、横線＝温度ごとの平均。"
                  "換算前は低温ほど上・高温ほど下にずれるが、温度シフトファクターで換算すると、ずれは約 4〜5 割に減る", 15, align="center", ls=1.3)
    y += h + 4
    row_label(yb_, y - 2, "温度シフトの効果")
    # (c) 予測検証
    yc = y
    pw = 176.0
    num = {"PP": ("13", "24", "6", "89"), "MIX": ("5", "6", "4", "98"), "GPPS": ("9", "13", "7", "93")}
    for i, (key, _) in enumerate(MATS):
        P.pic(fo(f"C_予測検証_{key}.png"), fx[i], y, w=pw)
        e, lo, hi, cv = num[key]
        nx, nw = fx[i] + pw + 2, fw - pw - 2
        t = (f"[s=13][c=52606D]平均的な誤差[/c][/s]\n[f={NUMF}][b][s=40][c={MCOL[key]}]{e}[/c][/s][/b][/f][s=18][b] %[/b][/s]\n"
             f"[s=13][c=52606D]低射出率側 {lo}%\n高射出率側 {hi}%[/c][/s]\n"
             f"[s=13][c=52606D]予測区間に入った割合[/c][/s]\n[f={NUMF}][b][s=34][c={MCOL[key]}]{cv}[/c][/s][/b][/f][s=18][b] %[/b][/s]")
        P.text(nx, y + 10, nw, t, 13, ls=1.25, after=3)
    y += 168 + 1
    _, h = P.text(fx[0], y, 3 * fw + 2 * gap, "学習に使っていない射出率条件（検証データ）の、実測と予測の粘度。点が破線に近いほど正確。縦線は 95% 予測区間。"
                  "誤差は、粘度の対数の二乗平均平方根誤差（RMSE）を % に換算した値", 15, align="center", ls=1.3)
    y += h + 4
    row_label(yc, y - 2, "予測検証")
    y += 14

    # ---------------- 3 段目：考察・まとめ
    ys = y
    # 5 どちらの温度でずらすか
    x = xs[0]
    y = heading(x, ys, cw, "5", "考察①：シフトに使う温度")
    ok, ng = "[s=26][c=2E8B57][b]○[/b][/c][/s]", "[s=26][c=D1495B][b]×[/b][/c][/s]"
    rows = [["シフトに使う温度", "PP", "粉砕材\n混合 PP", "GPPS"],
            ["金型内樹脂\n最高温度", ok, ok, ok],
            ["加熱シリンダ\n温度", ng, ng, ok + "※"]]
    yb = P.table(x, y, [80.0, (cw - 80) / 3, (cw - 80) / 3, (cw - 80) / 3], rows, 16, row_h=19.0, head_fill=CH, fills=("FFFFFF", LT),
                 line=LINE, lw=0.75)
    t = ("○：シフトファクターを組み込むと、データへの当てはまりが統計的に有意に向上（尤度比検定、有意水準 5%）。※GPPS は基準をわずかに上回る程度\n"
         "→ 設定温度（加熱シリンダ温度）よりも、[b]金型内で実際に測った樹脂温度[/b]でずらすほうが、データによく合う\n"
         "ただし、金型内樹脂最高温度は、せん断発熱で射出率とともに変わるため、このシフトファクターには[b]せん断発熱の影響も含まれる可能性[/b]がある")
    _, h = P.text(x, yb + 3, cw, t, 17, ls=1.4, after=4)
    b1 = yb + 3 + h
    # 6 従来の粘度式との比較
    x = xs[1]
    y = heading(x, ys, cw, "6", "考察②：従来の粘度式との比較")
    _, hb = P.pic(fo("C_誤差の比較.png"), x, y, w=cw)
    t = ("同じ学習データに当てはめた Cross-WLF モデルの予測誤差は約 59〜97% と、本モデル（約 5〜13%）より大幅に大きい。"
         "キャビティ厚みによる見かけの粘度の違いを表せないためと考えられる")
    _, h = P.text(x, y + hb + 1, cw, t, 15.5, ls=1.38)
    b2 = y + hb + 1 + h
    # 7 まとめ
    x = xs[2]
    y = heading(x, ys, cw, "7", "まとめと今後の課題")
    t = ("・金型内センサの計測データとガウス過程回帰で、樹脂粘度の[b]予測値と 95% 予測区間[/b]を同時に求めるモデルを、3 材料で構築した\n"
         "・学習に使っていない射出率条件で、誤差は約 5〜13%、実測値の 89〜98% が 95% 予測区間に入った\n"
         "・温度シフトファクターで、温度によるずれを約 4〜5 割に減らした\n"
         "・[b][c=D1495B]課題[/c][/b]：データのない約 150 1/s より小さい範囲では、予測区間が大きく広がり、曲線の形が不自然になることもある。"
         "低せん断速度側のデータの追加や、Cross 式などとの組み合わせを検討する")
    s7, h = P.box(x, y, cw, t, 18.5, fill=LT, radius=3, pad=(6, 6, 5, 5), ls=1.42, hang=1.0, after=5)
    b3 = y + h
    if max(b1, b2) > b3:                                             # まとめの枠を、隣の列の下端までのばす
        s7.height = int((max(b1, b2) - y) * 36000)
    y = max(b1, b2, b3) + 6
    _, h = P.text(M, y, CW, "【参考文献】" + REFS, 13, color=SUB, ls=1.2)
    return P, y + h


def main():
    which = [a.upper() for a in sys.argv[1:]] or ["A", "B", "C"]
    builders = {"A": (design_a, "オリジナル案A_ひと目でわかる版.pptx"), "B": (design_b, "オリジナル案B_QandAで読み解く版.pptx"),
                "C": (design_c, "オリジナル案C_材料別に比べる版.pptx")}
    for k in which:
        if k not in builders:
            print(f"案{k} はまだない")
            continue
        fn, name = builders[k]
        P, bottom = fn()
        out = os.path.join(POSTER, name)
        P.save(out, TITLE)
        print(f"案{k}：下端 {bottom:.1f} mm（ページの高さ {H:.0f} mm）{'　★はみ出し' if bottom > H else ''}　保存: {out}")
        for w in P.warnings:
            print("  ★", w)


if __name__ == "__main__":
    main()
