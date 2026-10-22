# -*- coding: utf-8 -*-
"""ポスター（PowerPoint）を python-pptx で組み立てる部品（build_poster_designs.py が使う）。

座標・大きさはすべて mm、文字の大きさは pt。文字は簡単な書式で書ける：
  [b]太字[/b]、[c=RRGGBB]色[/c]、[s=40]大きさ[/s]、[f=フォント名]書体[/f]、_{下付き}、^{上付き}（入れ子にできる）
文字の高さは、全角 1 字 ＝ 1 em、半角 ≒ 0.56 em として折り返しの行数を見積もって決める（実際より少し高めに出る）。
決めた高さに文字が収まらないと見積もったときは、warnings に記録する（build_poster_designs.py が最後に表示する）。
"""
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

PT_MM = 25.4 / 72.0
EMU_MM = 36000.0
TAG = re.compile(r"(\[/[bcsf]\]|\[[bcsf](?:=[^\]]*)?\]|_\{[^}]*\}|\^\{[^}]*\})")
OPEN = re.compile(r"\[([bcsf])(?:=([^\]]*))?\]")
SYM = re.compile("(γ̇)")          # γ̇（γ ＋ 上の点の結合文字）。和文の書体では点がずれるので、この字だけ数式用の書体で書く
MATH_FONT = "Cambria Math"
BULLETS = "・①②③④⑤⑥●■▶"
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def plain(text):
    """書式の記号を除いた文字"""
    return "".join(t for t, _ in segments(text, {"size": 10}))


def segments(text, base):
    """書式つきの文字 → [(文字, 書式の辞書), ...]"""
    stack, out = [dict(base)], []
    for tok in TAG.split(text):
        if not tok:
            continue
        if tok.startswith("[/"):
            if len(stack) > 1:
                stack.pop()
            continue
        m = OPEN.fullmatch(tok)
        if m:
            st, (k, v) = dict(stack[-1]), m.groups()
            if k == "b":
                st["bold"] = True
            elif k == "c":
                st["color"] = v
            elif k == "s":
                st["size"] = float(v)
            elif k == "f":
                st["font"] = v
            stack.append(st)
            continue
        st = dict(stack[-1])
        if tok.startswith("_{"):
            tok, st["baseline"] = tok[2:-1], -25000
        elif tok.startswith("^{"):
            tok, st["baseline"] = tok[2:-1], 30000
        out.append((tok, st))
    return out


def char_em(c):
    """1 字の幅の目安 [em]"""
    o = ord(c)
    if o in (0x0307,):
        return 0.0
    if o in (0x2013, 0x2212):
        return 0.6
    if o >= 0x2000:
        return 1.0
    if c == " ":
        return 0.3
    if c in "ilIjtf.,:;|!'()[]":
        return 0.36
    if c.isupper() or c in "mwMW%":
        return 0.74
    return 0.58


def text_width(text, base):
    """1 行で書いたときの幅の目安 [mm]"""
    return sum(char_em(c) * st["size"] * (0.7 if st.get("baseline") else 1.0) * PT_MM for t, st in segments(text, base) for c in t)


def para_lines(para, base, width):
    """段落の行数と、いちばん大きい文字 [pt]（全角の字は 1 字ずつ、半角の語は語の切れ目で折り返すとして数える）"""
    segs = segments(para, base)
    maxs = max([st["size"] for _, st in segs] or [base["size"]])
    toks, word = [], 0.0                              # 折り返しの単位：全角の字・空白は 1 字ずつ、半角の語はまとめて 1 つ
    for t, st in segs:
        f = (0.7 if st.get("baseline") else 1.0) * st["size"] * PT_MM
        for c in t:
            cw = char_em(c) * f
            if char_em(c) >= 1.0 or c == " ":
                if word:
                    toks.append(word)
                    word = 0.0
                toks.append(cw)
            else:
                word += cw
    if word:
        toks.append(word)
    n, x = 1, 0.0
    for tw in toks:
        if x + tw > width and x > 0:
            n, x = n + 1, 0.0
        if tw > width:
            n += int(tw // width)
            x = tw % width
        else:
            x += tw
    return n, maxs


class Poster:
    """1 枚のスライド（ポスター）。warnings：文字があふれそうな所の記録"""

    def __init__(self, w, h, font, color="000000"):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Mm(w), Mm(h)
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.W, self.H, self.font, self.color = w, h, font, color
        self.warnings = []

    # ------------------------------------------------------------ 文字
    def base(self, size, color=None, bold=False, font=None):
        return {"size": size, "color": color or self.color, "bold": bold, "font": font or self.font}

    def style_run(self, run, st, math=False):
        f = run.font
        f.size = Pt(st["size"] * (1.0 if not st.get("baseline") else 1.0))
        f.bold = bool(st.get("bold"))
        f.color.rgb = RGBColor.from_string(st["color"])
        font = MATH_FONT if math else st["font"]
        f.name = font
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", font)
        if st.get("baseline"):
            rPr.set("baseline", str(st["baseline"]))

    def text_height(self, text, base, width, ls=1.25, after=0.0, hang=None):
        h = 0.0
        paras = text.split("\n")
        for i, para in enumerate(paras):
            w = width - (hang * base["size"] * PT_MM if hang and plain(para)[:1] in BULLETS else 0.0)
            n, maxs = para_lines(para, base, w)
            h += n * maxs * ls * PT_MM + (after * PT_MM if i < len(paras) - 1 else 0.0)
        return h

    def fill_tf(self, tf, text, base, align="left", ls=1.25, after=0.0, hang=None):
        for i, para in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            segs = segments(para, base)
            maxs = max([st["size"] for _, st in segs] or [base["size"]])
            p.line_spacing = Pt(maxs * ls)
            p.space_after = Pt(after)
            p.space_before = Pt(0)
            p.font.size = Pt(base["size"])
            p.font.name = base["font"]
            if hang and plain(para)[:1] in BULLETS:
                pPr = p._p.get_or_add_pPr()
                w = int(hang * base["size"] * 12700)
                pPr.set("marL", str(w))
                pPr.set("indent", str(-w))
            for t, st in segs:
                for piece in SYM.split(t):
                    if piece:
                        r = p.add_run()
                        r.text = piece
                        self.style_run(r, st, math=bool(SYM.fullmatch(piece)))

    def frame(self, s, text, size, color=None, bold=False, font=None, align="left", anchor="top", ls=1.25, after=0.0, hang=None,
              margin=(0.0, 0.0, 0.0, 0.0), vert=False):
        """図形 s の中に文字を書く。margin：(左, 右, 上, 下) [mm]"""
        base = self.base(size, color, bold, font)
        tf = s.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ANCHOR[anchor]
        tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = (Mm(v) for v in margin)
        if vert:
            tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
        self.fill_tf(tf, text, base, align, ls, after, hang)
        return base

    def text(self, x, y, w, text, size, h=None, color=None, bold=False, font=None, align="left", anchor="top", ls=1.25, after=0.0, hang=None,
             margin=(0.0, 0.0, 0.0, 0.0), name=None):
        """文字の枠。h を省くと、見積もった高さにする。返り値：(図形, 高さ [mm])"""
        base = self.base(size, color, bold, font)
        need = self.text_height(text, base, w - margin[0] - margin[1], ls, after, hang) + margin[2] + margin[3]
        hh = need if h is None else h
        s = self.s.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(hh))
        self.frame(s, text, size, color, bold, font, align, anchor, ls, after, hang, margin)
        if h is not None and need > h + 0.8:
            self.warnings.append(f"文字があふれそう（{name or plain(text)[:24]}…）：高さ {h:.1f} mm に、見積もり {need:.1f} mm")
        return s, hh

    # ------------------------------------------------------------ 図形
    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, kind=None, alpha=None, shadow=None, dash=None):
        """四角（radius [mm] を指定すると角を丸める）。shadow：(ぼかし, ずれ, 濃さ) [mm, mm, 0〜1]"""
        kind = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        s = self.s.shapes.add_shape(kind, Mm(x), Mm(y), Mm(w), Mm(h))
        s.shadow.inherit = False
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
            if alpha is not None:
                clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(lw)
            if dash:
                ln = s.line._get_or_add_ln()
                etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
        else:
            s.line.fill.background()
        if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = min(0.5, radius / min(w, h))
        if shadow:
            self.add_shadow(s, *shadow)
        return s

    def add_shadow(self, s, blur=3.0, dist=1.2, alpha=0.18):
        spPr = s._element.spPr
        eff = spPr.find(qn("a:effectLst"))                 # shadow.inherit = False で作られた空の要素を使う（2 つあると、ファイルが壊れる）
        if eff is None:
            eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * EMU_MM)), dist=str(int(dist * EMU_MM)), dir="5400000",
                              algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))

    def box(self, x, y, w, text, size, h=None, fill="FFFFFF", line=None, lw=1.0, radius=None, pad=(3.0, 3.0, 2.5, 2.5), color=None,
            bold=False, font=None, align="left", anchor="top", ls=1.25, after=0.0, hang=None, shadow=None, kind=None, name=None, alpha=None):
        """文字の入った四角。h を省くと、文字に合わせた高さにする。返り値：(図形, 高さ)"""
        base = self.base(size, color, bold, font)
        need = self.text_height(text, base, w - pad[0] - pad[1], ls, after, hang) + pad[2] + pad[3]
        hh = need if h is None else h
        s = self.rect(x, y, w, hh, fill=fill, line=line, lw=lw, radius=radius, shadow=shadow, kind=kind, alpha=alpha)
        self.frame(s, text, size, color, bold, font, align, anchor, ls, after, hang, pad)
        if h is not None and need > h + 0.8:
            self.warnings.append(f"文字があふれそう（{name or plain(text)[:24]}…）：高さ {h:.1f} mm に、見積もり {need:.1f} mm")
        return s, hh

    def pill_w(self, text, size, padx=4.0, bold=True, font=None):
        """札（pill）の幅 [mm]（文字の幅の見積もりは少し小さめに出るので、6% 足す）"""
        return text_width(text, self.base(size, None, bold, font)) * 1.06 + 2 * padx

    def pill(self, x, y, text, size, fill, color="FFFFFF", h=None, padx=4.0, bold=True, font=None, align="left", w=None):
        """文字の長さに合わせた、角の丸い札（文字は折り返さない）。align="right" なら x を右端とみなす。返り値：(図形, 幅, 高さ)"""
        hh = h or size * PT_MM * 1.55
        ww = w or self.pill_w(text, size, padx, bold, font)
        xx = x - ww if align == "right" else (x - ww / 2 if align == "center" else x)
        s, _ = self.box(xx, y, ww, text, size, h=hh, fill=fill, radius=hh / 2, pad=(0.5, 0.5, 0.0, 0.0), color=color, bold=bold, font=font,
                        align="center", anchor="middle", ls=1.0)
        s.text_frame.word_wrap = False
        return s, ww, hh

    def circle(self, cx, cy, d, text="", size=20, fill="000000", color="FFFFFF", bold=True, font=None, line=None, lw=1.0):
        s = self.rect(cx - d / 2, cy - d / 2, d, d, fill=fill, line=line, lw=lw, kind=MSO_SHAPE.OVAL)
        if text:
            self.frame(s, text, size, color, bold, font, "center", "middle", ls=1.0)
        return s

    def shape(self, kind, x, y, w, h, fill=None, line=None, lw=1.0, rotation=0.0, adj=None):
        s = self.rect(x, y, w, h, fill=fill, line=line, lw=lw, kind=kind)
        if rotation:
            s.rotation = rotation
        if adj:
            for i, v in enumerate(adj):
                s.adjustments[i] = v
        return s

    def line(self, x1, y1, x2, y2, color="000000", lw=1.0, dash=None, head=None, tail=None):
        """直線。head・tail：矢じりの形（"triangle" など）。head が終点側"""
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
        c.line.color.rgb = RGBColor.from_string(color)
        c.line.width = Pt(lw)
        ln = c.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
        if tail:
            etree.SubElement(ln, qn("a:headEnd"), type=tail, w="med", len="med")
        if head:
            etree.SubElement(ln, qn("a:tailEnd"), type=head, w="med", len="med")
        return c

    def pic(self, path, x, y, w=None, h=None):
        p = self.s.shapes.add_picture(path, Mm(x), Mm(y), Mm(w) if w else None, Mm(h) if h else None)
        return p.width / EMU_MM, p.height / EMU_MM

    def table(self, x, y, widths, rows, size, row_h=None, head_fill="333333", head_color="FFFFFF", fills=("FFFFFF", "F2F2F2"), first_fill=None,
              line="FFFFFF", lw=1.0, aligns=None, bold_first=True, color=None, pad=1.8, head_size=None):
        """文字の入った四角を並べた表（1 行目は見出し）。返り値：下端 [mm]"""
        rh = row_h or size * PT_MM * 1.9
        yy = y
        for i, row in enumerate(rows):
            xx = x
            for j, (cw, t) in enumerate(zip(widths, row)):
                head = i == 0
                fill = head_fill if head else (first_fill if (j == 0 and first_fill) else fills[(i - 1) % len(fills)])
                al = "center" if head else (aligns[j] if aligns else ("left" if j == 0 else "center"))
                self.box(xx, yy, cw, t, head_size if head and head_size else size, h=rh, fill=fill, line=line, lw=lw, pad=(pad, pad, 0, 0),
                         color=head_color if head else color, bold=head or (bold_first and j == 0), align=al, anchor="middle", ls=1.0)
                xx += cw
            yy += rh
        return yy

    def mark(self):
        """今の図形の数（to_back で、この位置の後ろに図形を回すために使う）"""
        return len(self.s.shapes)

    def to_back(self, shape, mark):
        """shape を、mark の時点の図形の直後（＝それより後に置いた図形の後ろ）に回す"""
        tree = self.s.shapes._spTree
        tree.remove(shape._element)
        tree.insert(2 + mark, shape._element)

    def save(self, path, title=""):
        self.prs.core_properties.title = title
        self.prs.save(path)
